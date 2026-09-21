import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import pymupdf
import pytest

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
PDF = FIXTURES / "pdf"


def run_openreview(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "openreview_cli", "parse", *args],
        capture_output=True,
        text=True,
    )


def _write_pdf(path: Path, author: str, title: str) -> None:
    doc: Any = pymupdf.open()  # type: ignore[no-untyped-call]
    page = doc.new_page()
    page.insert_text((72, 72), "Mutual NDA between Jane Doe and Acme Corp.")
    doc.set_metadata({"author": author, "title": title})
    doc.save(str(path))
    doc.close()


class TestParseCommand:
    @pytest.mark.integration
    def test_parse_without_path_shows_help(self) -> None:
        result = run_openreview()
        assert result.returncode != 0

    @pytest.mark.integration
    def test_parse_simple_contract(self) -> None:
        result = run_openreview(str(PDF / "simple_contract.pdf"))
        assert result.returncode == 0
        assert "clause-" in result.stdout

    @pytest.mark.integration
    def test_parse_json_output(self) -> None:
        result = run_openreview("--format", "json", str(PDF / "simple_contract.pdf"))
        assert result.returncode == 0
        data = json.loads(result.stdout)
        assert isinstance(data, list)
        assert len(data) > 0
        assert "id" in data[0]
        assert "text" in data[0]

    @pytest.mark.integration
    def test_parse_summary(self) -> None:
        result = run_openreview("--summary", str(PDF / "simple_contract.pdf"))
        assert result.returncode == 0
        assert "Parsed" in result.stdout
        assert "clauses" in result.stdout

    @pytest.mark.integration
    def test_parse_non_existent_file(self) -> None:
        result = run_openreview(str(FIXTURES / "nonexistent.pdf"))
        assert result.returncode == 8
        assert "No file found" in result.stderr

    @pytest.mark.integration
    def test_parse_unsupported_format(self) -> None:
        result = run_openreview(str(FIXTURES / "test.txt"))
        assert result.returncode == 8
        assert "supported" in result.stderr.lower()


class TestParseDocumentMetadata:
    @pytest.mark.integration
    def test_parse_document_populates_docx_metadata(self, tmp_path: Path) -> None:
        from docx import Document

        from openreview_cli.parsing.stream import parse_document

        path = tmp_path / "meta.docx"
        source = Document()
        source.add_paragraph("Mutual NDA between Jane Doe and Acme Corp.")
        source.core_properties.author = "Jane Doe"
        source.core_properties.title = "Mutual NDA"
        source.save(str(path))

        doc, clauses = parse_document(path)

        assert doc.author == "Jane Doe"
        assert doc.title == "Mutual NDA"
        assert doc.company is None
        assert clauses

    @pytest.mark.integration
    def test_parse_document_populates_pdf_metadata(self, tmp_path: Path) -> None:
        from openreview_cli.parsing.stream import parse_document

        path = tmp_path / "meta.pdf"
        _write_pdf(path, "Jane Doe", "Mutual NDA")

        doc, clauses = parse_document(path)

        assert doc.author == "Jane Doe"
        assert doc.title == "Mutual NDA"
        assert doc.company is None
        assert clauses

    @pytest.mark.integration
    def test_parse_document_survives_metadata_failure(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from openreview_cli.parsing.stream import parse_document

        path = tmp_path / "meta.pdf"
        _write_pdf(path, "Jane Doe", "Mutual NDA")

        def _boom(_doc: object) -> object:
            raise RuntimeError("boom")

        monkeypatch.setattr("openreview_cli.parsing.pdf_parser.pdf_metadata", _boom)

        doc, clauses = parse_document(path)

        assert doc.author is None
        assert doc.title is None
        assert clauses


class TestParseCommandWarnings:
    @pytest.mark.integration
    def test_parse_shows_non_english_warning_on_stderr(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from docx import Document

        from openreview_cli.app import parse

        path = tmp_path / "arabic.docx"
        source = Document()
        source.add_paragraph("مرحبا بالعالم")
        source.add_paragraph("Hello world")
        source.save(str(path))

        parse(str(path), "text", False)

        captured = capsys.readouterr()
        assert "The contract appears to be in Arabic. Results may be less accurate" in captured.err

    @pytest.mark.integration
    def test_parse_clean_document_prints_no_warnings(
        self, tmp_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from docx import Document

        from openreview_cli.app import parse

        path = tmp_path / "clean.docx"
        source = Document()
        source.add_paragraph("This agreement is governed by the laws of Delaware.")
        source.save(str(path))

        parse(str(path), "text", False)

        captured = capsys.readouterr()
        assert captured.err == ""
        assert "Results may be less accurate" not in captured.out
        assert "clause-" in captured.out
