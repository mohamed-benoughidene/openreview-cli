from pathlib import Path

import pytest

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
PDF = FIXTURES / "pdf"
DOCX = FIXTURES / "docx"


class TestStreamClauses:
    @pytest.mark.integration
    def test_routes_to_pdf_parser_for_pdf(self) -> None:
        from openreview_cli.parsing.stream import stream_clauses

        clauses = list(stream_clauses(PDF / "simple_contract.pdf"))
        assert all(c.source_page is not None for c in clauses)

    @pytest.mark.integration
    def test_routes_to_docx_parser_for_docx(self) -> None:
        from openreview_cli.parsing.stream import stream_clauses

        clauses = list(stream_clauses(DOCX / "simple_contract.docx"))
        assert all(c.source_paragraph is not None for c in clauses)

    @pytest.mark.integration
    def test_unsupported_format_raises(self) -> None:
        from openreview_cli.parsing.models import ParseError
        from openreview_cli.parsing.stream import stream_clauses

        with pytest.raises(ParseError) as exc:
            list(stream_clauses(FIXTURES / "test.txt"))
        assert exc.value.category == "unsupported_format"

    @pytest.mark.integration
    def test_non_existent_path_raises(self) -> None:
        from openreview_cli.parsing.models import ParseError
        from openreview_cli.parsing.stream import stream_clauses

        with pytest.raises(ParseError) as exc:
            list(stream_clauses(FIXTURES / "nonexistent.pdf"))
        assert exc.value.category == "file_not_found"

    @pytest.mark.integration
    def test_cross_format_equivalence_simple(self) -> None:
        from openreview_cli.parsing.stream import stream_clauses

        pdf_clauses = list(stream_clauses(PDF / "simple_contract.pdf"))
        docx_clauses = list(stream_clauses(DOCX / "simple_contract.docx"))
        pdf_count = len(pdf_clauses)
        docx_count = len(docx_clauses)
        diff = abs(pdf_count - docx_count) / max(pdf_count, docx_count)
        assert diff <= 0.1

    @pytest.mark.integration
    def test_cross_format_equivalence_flat(self) -> None:
        from openreview_cli.parsing.stream import stream_clauses

        pdf_clauses = list(stream_clauses(PDF / "flat_document.pdf"))
        docx_clauses = list(stream_clauses(DOCX / "flat_document.docx"))
        assert all(c.level == 0 for c in pdf_clauses)
        assert all(c.level == 0 for c in docx_clauses)


class TestCrossFormatHierarchy:
    @pytest.mark.integration
    def test_parent_id_chain_integrity(self) -> None:
        from openreview_cli.parsing.stream import stream_clauses

        for path in [PDF / "simple_contract.pdf", DOCX / "simple_contract.docx"]:
            clauses = list(stream_clauses(path))
            ids = {c.id for c in clauses}
            for clause in clauses:
                if clause.parent_id is not None:
                    assert clause.parent_id in ids

    @pytest.mark.integration
    def test_english_fixtures_have_no_warnings(self) -> None:
        from openreview_cli.parsing.stream import parse_document

        for path in (PDF / "simple_contract.pdf", DOCX / "simple_contract.docx"):
            doc, clauses = parse_document(path)
            assert doc.warnings == []
            assert all(not c.is_non_english for c in clauses)

    @pytest.mark.integration
    def test_non_english_docx_is_flagged_end_to_end(self, tmp_path: Path) -> None:
        from docx import Document as DocxDocument

        from openreview_cli.parsing.stream import parse_document

        docx_path = tmp_path / "arabic.docx"
        source = DocxDocument()
        source.add_paragraph("مرحبا بالعالم")
        source.add_paragraph("Hello world")
        source.save(str(docx_path))

        doc, clauses = parse_document(docx_path)

        assert doc.warnings == [
            "The contract appears to be in Arabic. Results may be less accurate"
        ]
        assert any(c.is_non_english is True for c in clauses)

    @pytest.mark.integration
    def test_tofu_docx_is_flagged_end_to_end(self, tmp_path: Path) -> None:
        from docx import Document as DocxDocument

        from openreview_cli.parsing.stream import parse_document

        docx_path = tmp_path / "tofu.docx"
        source = DocxDocument()
        source.add_paragraph("Broken text \ufffd here")
        source.save(str(docx_path))

        doc, _clauses = parse_document(docx_path)

        assert "Some text could not be read correctly. Results may contain errors" in doc.warnings
