"""Unit tests for PDF/DOCX metadata extraction (Gap #8)."""

import re
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pymupdf
from docx import Document

from openreview_cli.parsing.docx_parser import DocxParser, docx_metadata
from openreview_cli.parsing.pdf_parser import PdfParser, pdf_metadata

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
PDF_FIXTURES = FIXTURES / "pdf"


def _write_pdf(path: Path, author: str, title: str) -> None:
    doc: Any = pymupdf.open()  # type: ignore[no-untyped-call]
    page = doc.new_page()
    page.insert_text((72, 72), "Mutual NDA between Jane Doe and Acme Corp.")
    doc.set_metadata({"author": author, "title": title})
    doc.save(str(path))
    doc.close()


def _write_docx(path: Path, author: str, title: str) -> None:
    doc = Document()
    doc.add_paragraph("Mutual NDA between Jane Doe and Acme Corp.")
    doc.core_properties.author = author
    doc.core_properties.title = title
    doc.save(str(path))


def _rewrite_app_xml(source: Path, dest: Path, company: str | None) -> None:
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == "docProps/app.xml":
                if company is None:
                    continue
                text = zin.read(item.filename).decode("utf-8")
                text = re.sub(
                    r"<Company\s*/>|<Company>.*?</Company>",
                    f"<Company>{company}</Company>",
                    text,
                )
                zout.writestr(item, text.encode("utf-8"))
                continue
            data = zin.read(item.filename)
            if item.filename == "_rels/.rels" and company is None:
                rels = data.decode("utf-8")
                rels = re.sub(r'<Relationship[^>]*Target="docProps/app\.xml"[^>]*/>', "", rels)
                data = rels.encode("utf-8")
            zout.writestr(item, data)


def test_pdf_metadata_reads_author_and_title(tmp_path: Path) -> None:
    path = tmp_path / "meta.pdf"
    _write_pdf(path, "Jane Doe", "Mutual NDA")

    with pymupdf.open(str(path)) as doc:  # type: ignore[no-untyped-call]
        assert pdf_metadata(doc) == ("Jane Doe", "Mutual NDA")


def test_pdf_metadata_empty_strings_become_none() -> None:
    with pymupdf.open(str(PDF_FIXTURES / "simple_contract.pdf")) as doc:  # type: ignore[no-untyped-call]
        assert pdf_metadata(doc) == (None, None)


def test_docx_metadata_reads_author_and_title(tmp_path: Path) -> None:
    path = tmp_path / "meta.docx"
    _write_docx(path, "Jane Doe", "Mutual NDA")

    assert docx_metadata(Document(str(path)))[:2] == ("Jane Doe", "Mutual NDA")


def test_docx_metadata_reads_company_from_app_xml(tmp_path: Path) -> None:
    base = tmp_path / "base.docx"
    _write_docx(base, "Jane Doe", "Mutual NDA")
    out = tmp_path / "company.docx"
    _rewrite_app_xml(base, out, "Acme Corp")

    assert docx_metadata(Document(str(out)))[2] == "Acme Corp"


def test_docx_metadata_without_app_xml_returns_none_company(tmp_path: Path) -> None:
    base = tmp_path / "base.docx"
    _write_docx(base, "Jane Doe", "Mutual NDA")
    out = tmp_path / "noapp.docx"
    _rewrite_app_xml(base, out, None)

    assert docx_metadata(Document(str(out)))[2] is None


def test_pdf_metadata_handles_missing_metadata_dict() -> None:
    assert pdf_metadata(SimpleNamespace()) == (None, None)


def test_parser_sets_metadata_attributes(tmp_path: Path) -> None:
    pdf_path = tmp_path / "meta.pdf"
    _write_pdf(pdf_path, "Jane Doe", "Mutual NDA")
    pdf_parser = PdfParser(pdf_path)
    assert len(list(pdf_parser.parse())) > 0
    assert pdf_parser.author == "Jane Doe"
    assert pdf_parser.title == "Mutual NDA"
    assert pdf_parser.company is None

    docx_path = tmp_path / "meta.docx"
    _write_docx(docx_path, "Jane Doe", "Mutual NDA")
    docx_parser = DocxParser(docx_path)
    assert len(list(docx_parser.parse())) > 0
    assert docx_parser.author == "Jane Doe"
    assert docx_parser.title == "Mutual NDA"
    assert docx_parser.company is None
