"""Unit tests for PII stripping engine."""

import time as time_mod
from pathlib import Path
from unittest.mock import patch

from _pytest.logging import LogCaptureFixture
from _pytest.monkeypatch import MonkeyPatch

from openreview_cli.parsing.models import Clause, Document
from openreview_cli.pii.engine import PiiEngine, strip_pii, strip_pii_clauses
from openreview_cli.pii.models import PiiEntity


class TestStripPii:
    """Tests for the strip_pii top-level function."""

    def test_strip_skipped_when_disabled(self) -> None:
        """When PII stripping is disabled, no entities are detected and a warning is emitted."""
        clause = Clause(
            id="1",
            title=None,
            text="John Smith works at ABC Corp.",
            level=1,
            parent_id=None,
            source_page=1,
            source_paragraph=None,
            source_span=None,
        )
        doc = Document(
            source_path=Path("/tmp/test.pdf"),
            format="pdf",
            page_count=1,
            clause_count=1,
            parse_duration_seconds=0.1,
            warnings=[],
        )

        engine = PiiEngine(threshold=0.7)
        with patch.object(
            engine, "detect_all_pages", return_value=([], ["PII detection disabled"], [], {})
        ):
            result = strip_pii([clause], doc, strip_metadata=False, engine=engine)

        assert any("disabled" in w for w in result.warnings)
        assert result.mapping == {}
        assert result.stripped_text == "John Smith works at ABC Corp."


class TestStripPiiClauses:
    """Tests for the strip_pii_clauses top-level function."""

    def make_clause(
        self,
        id: str = "1",
        text: str = "Test text",
        title: str | None = "Test Clause",
        level: int = 1,
        parent_id: str | None = None,
        source_page: int | None = 1,
        source_paragraph: int | None = None,
        source_span: tuple[int, int] | None = None,
    ) -> Clause:
        return Clause(
            id=id,
            title=title,
            text=text,
            level=level,
            parent_id=parent_id,
            source_page=source_page,
            source_paragraph=source_paragraph,
            source_span=source_span,
        )

    def make_doc(self, source_path: str = "/tmp/test.pdf") -> Document:
        return Document(
            source_path=Path(source_path),
            format="pdf",
            page_count=1,
            clause_count=1,
            parse_duration_seconds=0.1,
            warnings=[],
        )

    def test_preserves_metadata(self) -> None:
        """T001: All clause metadata fields unchanged after stripping."""
        clause = self.make_clause(
            id="clause-42",
            text="Acme Corp shall pay $1,000.",
            title="Payment Terms",
            level=2,
            parent_id="clause-1",
            source_page=3,
            source_paragraph=None,
            source_span=(10, 50),
        )
        doc = self.make_doc()
        engine = PiiEngine(threshold=0.7)
        entity = PiiEntity(
            entity_type="ORGANIZATION",
            original_value="Acme Corp",
            start=0,
            end=9,
            score=0.9,
            placeholder="[TEMP_0]",
            source="nlp",
        )
        with patch.object(engine, "detect_all_pages", return_value=([entity], [], [], {})):
            stripped, _result = strip_pii_clauses(
                [clause], doc, strip_metadata=False, engine=engine
            )

        out = stripped[0]
        assert out.id == "clause-42"
        assert out.title == "Payment Terms"
        assert out.level == 2
        assert out.parent_id == "clause-1"
        assert out.source_page == 3
        assert out.source_paragraph is None
        assert out.source_span == (10, 50)

    def test_replaces_pii(self) -> None:
        """T002: Clause text contains placeholders instead of raw PII."""
        clause = self.make_clause(text="John Smith works at Acme Corp.")
        doc = self.make_doc()
        engine = PiiEngine(threshold=0.7)
        entities = [
            PiiEntity(
                entity_type="PERSON",
                original_value="John Smith",
                start=0,
                end=10,
                score=0.95,
                placeholder="[TEMP_0]",
                source="nlp",
            ),
            PiiEntity(
                entity_type="ORGANIZATION",
                original_value="Acme Corp",
                start=20,
                end=29,
                score=0.9,
                placeholder="[TEMP_0]",
                source="nlp",
            ),
        ]
        with patch.object(engine, "detect_all_pages", return_value=(entities, [], [], {})):
            stripped, result = strip_pii_clauses([clause], doc, strip_metadata=False, engine=engine)

        assert "John Smith" not in stripped[0].text
        assert "Acme Corp" not in stripped[0].text
        assert "[NAME_1]" in stripped[0].text or "[PARTY_1]" in stripped[0].text

        assert result.mapping is not None

    def test_empty_input(self) -> None:
        """T003: Empty list returns ([], PiiResult) with empty mapping."""
        doc = self.make_doc()
        engine = PiiEngine(threshold=0.7)
        with patch.object(engine, "detect_all_pages", return_value=([], [], [], {})):
            stripped, result = strip_pii_clauses([], doc, strip_metadata=False, engine=engine)

        assert stripped == []
        assert result.mapping == {}

    def test_no_pii_unchanged(self) -> None:
        """T004: Clause with no PII has unchanged text and no mapping entries."""
        clause = self.make_clause(text="This clause has no sensitive data.")
        doc = self.make_doc()
        engine = PiiEngine(threshold=0.7)
        with patch.object(engine, "detect_all_pages", return_value=([], [], [], {})):
            stripped, result = strip_pii_clauses([clause], doc, strip_metadata=False, engine=engine)

        assert stripped[0].text == "This clause has no sensitive data."
        assert result.mapping == {}

    def test_metadata_entities(self) -> None:
        """T005: Metadata placeholders appended to last clause when not in clause text."""
        clause_a = self.make_clause(id="1", text="Clause one.", source_page=1)
        clause_b = self.make_clause(id="2", text="Clause two.", source_page=1)
        doc = Document(
            source_path=Path("/tmp/report.pdf"),
            format="pdf",
            page_count=1,
            clause_count=2,
            parse_duration_seconds=0.1,
            warnings=[],
            author="Legal Dept",
            title="Annual Report",
        )
        engine = PiiEngine(threshold=0.7)
        with patch.object(engine, "detect_all_pages", return_value=([], [], [], {})):
            stripped, result = strip_pii_clauses([clause_a, clause_b], doc, engine=engine)

        metadata_placeholders = [
            k
            for k in result.mapping
            if result.mapping[k] in ("report.pdf", "Legal Dept", "Annual Report")
        ]
        assert len(metadata_placeholders) > 0, "Expected at least one metadata placeholder"

        # Last clause should have metadata placeholders appended
        for key in metadata_placeholders:
            assert f"[{key}]" in stripped[-1].text, (
                f"Metadata placeholder [{key}] not found in last clause text"
            )

    def test_performance(self) -> None:
        """T006: pathology guard — strip_pii_clauses must not regress super-linearly vs strip_pii.

        This is a sanity guard, not a performance contract. Warm, interleaved, best-of-5 timing
        of both functions over the same 200-clause document is stable at ~3.5x locally; the bound
        below (15.0x) leaves ~4x headroom over that while still catching a pathological regression
        such as an O(n^2) wrapper. The previous shape timed strip_pii_clauses cold against a warm
        strip_pii, so a loaded 2-vCPU CI runner measured 18.72x for unchanged code.
        """
        # 200 clauses to amortize per-clause overhead relative to mock baseline
        clauses = [
            self.make_clause(id=str(i), text=f"Clause {i} with Acme Corp data.", source_page=1)
            for i in range(200)
        ]
        doc = self.make_doc()
        engine = PiiEngine(threshold=0.7)
        entity = PiiEntity(
            entity_type="ORGANIZATION",
            original_value="Acme Corp",
            start=0,
            end=9,
            score=0.9,
            placeholder="[TEMP_0]",
            source="nlp",
        )

        with patch.object(engine, "detect_all_pages", return_value=([entity], [], [], {})):
            # Warm both functions once so neither pays first-call costs inside the timer.
            strip_pii_clauses(clauses, doc, strip_metadata=False, engine=engine)
            strip_pii(clauses, doc, strip_metadata=False, engine=engine)

            # Interleaved best-of-5: alternate the two functions so a load spike cannot land on
            # only one of them. Noise only ever adds time, so the minimum is the stable estimator.
            t_bridge = float("inf")
            t_strip = float("inf")
            for _ in range(5):
                t0 = time_mod.perf_counter()
                strip_pii_clauses(clauses, doc, strip_metadata=False, engine=engine)
                t_bridge = min(t_bridge, time_mod.perf_counter() - t0)

                t0 = time_mod.perf_counter()
                strip_pii(clauses, doc, strip_metadata=False, engine=engine)
                t_strip = min(t_strip, time_mod.perf_counter() - t0)

        ratio = t_bridge / max(t_strip, 1e-9)
        # ponytail: sanity guard, not a perf contract — warm best-of-5 baseline is ~3.5x locally;
        # the old cold-vs-warm shape measured 18.72x on CI for unchanged code.
        assert ratio < 15.0, f"strip_pii_clauses {ratio:.2f}x slower than strip_pii (limit: 15.0x)"


class TestPiiEngineIsAvailable:
    """T006: PiiEngine.is_available() readiness check."""

    def test_returns_true_when_engine_ready(self) -> None:
        engine = PiiEngine(threshold=0.7)
        with patch.object(engine, "_ensure_analyzer", return_value=_MockAnalyzer()):
            assert engine.is_available() is True

    def test_returns_false_on_engine_failure(self) -> None:
        engine = PiiEngine(threshold=0.7)

        class FailingAnalyzer:
            def analyze(self, **_kw: object) -> object:
                msg = "model not found"
                raise RuntimeError(msg)

        with patch.object(engine, "_ensure_analyzer", return_value=FailingAnalyzer()):
            assert engine.is_available() is False

    def test_caches_result(self) -> None:
        engine = PiiEngine(threshold=0.7)
        mock = _MockAnalyzer()
        with patch.object(engine, "_ensure_analyzer", return_value=mock):
            assert engine.is_available() is True
            assert engine.is_available() is True  # second call uses cache
        # After patching removed, cache should still hold
        assert engine.is_available() is True


class _MockAnalyzer:
    """Minimal mock for Presidio AnalyzerEngine."""

    def analyze(self, **_kw: object) -> list[object]:
        return []


def test_is_available_logs_warning_on_failure(
    pii_engine: PiiEngine,
    monkeypatch: MonkeyPatch,
    caplog: LogCaptureFixture,
) -> None:
    pii_engine._is_available_cache = None
    monkeypatch.setattr(
        pii_engine,
        "_ensure_analyzer",
        lambda: (_ for _ in ()).throw(RuntimeError("boom")),
    )
    with caplog.at_level("WARNING", logger="openreview_cli.pii.engine"):
        assert pii_engine.is_available() is False
    assert "boom" in caplog.text
    pii_engine._is_available_cache = None


def test_annotate_clauses_output_reaches_pii_engine_flagged(
    pii_engine: PiiEngine, monkeypatch: MonkeyPatch
) -> None:
    from openreview_cli.parsing.clause_detector import annotate_clauses

    clauses = [
        Clause(
            id="clause-1",
            title="Non-English",
            text="مرحبا بالعالم",
            level=0,
            parent_id=None,
            source_page=1,
            source_paragraph=None,
            source_span=None,
        )
    ]
    annotate_clauses(clauses)
    assert clauses[0].is_non_english is True

    captured: dict[str, bool] = {}

    def _record(text: str, **kwargs: object) -> list[object]:
        captured["is_non_english"] = bool(kwargs.get("is_non_english", False))
        return []

    monkeypatch.setattr(pii_engine, "detect_on_page", _record)
    _entities, warnings, _failed_pages, _errors = pii_engine.detect_all_pages(clauses)

    assert captured["is_non_english"] is True
    assert any("Non-English" in w for w in warnings)


def test_detect_all_pages_emits_progress_via_callback(
    pii_engine: PiiEngine, monkeypatch: MonkeyPatch
) -> None:
    monkeypatch.setattr(pii_engine, "detect_on_page", lambda text, **kwargs: [])
    events: list[tuple[str, int, int]] = []
    from openreview_cli.parsing.models import Clause

    clauses = [
        Clause(
            id="1",
            title="C1",
            text="Contact john@example.com for details.",
            level=1,
            parent_id=None,
            source_page=1,
            source_paragraph=None,
            source_span=None,
        ),
        Clause(
            id="2",
            title="C2",
            text="Call Jane at 555-1234.",
            level=1,
            parent_id=None,
            source_page=2,
            source_paragraph=None,
            source_span=None,
        ),
    ]
    pii_engine.detect_all_pages(
        clauses,
        progress_callback=lambda desc, done, total: events.append((desc, done, total)),
    )
    assert events, "callback never invoked"
    assert events[-1][1] <= events[-1][2]


class TestDetectOnPageAllowlist:
    """detect_on_page restricts analysis to the 11 ground-truth entity types."""

    def test_analyze_called_with_entities_allowlist(self, monkeypatch: MonkeyPatch) -> None:
        engine = PiiEngine(threshold=0.7)
        captured: dict[str, object] = {}

        class _RecordingAnalyzer:
            def analyze(self, **kwargs: object) -> list[object]:
                captured.update(kwargs)
                return []

        monkeypatch.setattr(engine, "_ensure_analyzer", _RecordingAnalyzer)
        engine.detect_on_page("Some contract text with $5,000.")

        assert captured["entities"] == [
            "PERSON",
            "ORGANIZATION",
            "LOCATION",
            "DATE_TIME",
            "EMAIL_ADDRESS",
            "PHONE_NUMBER",
            "AMOUNT",
            "TAX_ID",
            "ACCT",
            "ID_DOCUMENT",
            "REG_NUMBER",
        ]


def test_metadata_entities_include_author_title_company() -> None:
    clause = Clause(
        id="1",
        title=None,
        text="Confidential terms apply.",
        level=1,
        parent_id=None,
        source_page=1,
        source_paragraph=None,
        source_span=None,
    )
    doc = Document(
        source_path=Path("/tmp/report.pdf"),
        format="pdf",
        page_count=1,
        clause_count=1,
        parse_duration_seconds=0.1,
        warnings=[],
        author="Jane Doe",
        title="Mutual NDA",
        company="Acme Corp",
    )
    engine = PiiEngine(threshold=0.7)
    with patch.object(engine, "detect_all_pages", return_value=([], [], [], {})):
        result = strip_pii([clause], doc, strip_metadata=True, engine=engine)

    assert result.mapping["AUTHOR_1"] == "Jane Doe"
    assert result.mapping["TITLE_1"] == "Mutual NDA"
    assert result.mapping["COMPANY_1"] == "Acme Corp"
    assert result.mapping["FILENAME_1"] == "report.pdf"
    assert "[AUTHOR_1]" in result.stripped_text
    assert "[TITLE_1]" in result.stripped_text
    assert "[COMPANY_1]" in result.stripped_text
    assert "[FILENAME_1]" in result.stripped_text


def test_python_docx_author_is_redacted_faithfully(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    from openreview_cli.parsing.stream import parse_document

    path = tmp_path / "generated.docx"
    source = DocxDocument()
    source.add_paragraph("Mutual NDA between Jane Doe and Acme Corp.")
    source.save(str(path))

    doc, clauses = parse_document(path)
    assert doc.author == "python-docx"

    engine = PiiEngine(threshold=0.7)
    with patch.object(engine, "detect_all_pages", return_value=([], [], [], {})):
        result = strip_pii(clauses, doc, strip_metadata=True, engine=engine)

    assert result.mapping["AUTHOR_1"] == "python-docx"
    assert "[AUTHOR_1]" in result.stripped_text
